"""Document export service: TXT, DOCX, PDF from edited content."""
import io
import re

ALLOWED_FORMATS = ("txt", "docx", "pdf")


def safe_filename(document_type: str, ext: str) -> str:
    ext = (ext or "").lower().strip()
    if ext not in ALLOWED_FORMATS:
        raise ValueError(f"Unsupported format: {ext}")
    base = (document_type or "document").lower().strip()
    base = re.sub(r"[^a-z0-9]+", "_", base).strip("_") or "document"
    return f"{base}.{ext}"


def export_txt(content: str) -> bytes:
    return (content or "").encode("utf-8")


def export_docx(content: str) -> bytes:
    from docx import Document

    doc = Document()
    for line in (content or "").splitlines() or [""]:
        if line.strip() == "":
            doc.add_paragraph("")
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        else:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_pdf(content: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    buf = io.BytesIO()
    template = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=1 * inch,
        rightMargin=1 * inch,
        topMargin=1 * inch,
        bottomMargin=1 * inch,
        title="LegalEase Document",
    )
    styles = getSampleStyleSheet()
    body = styles["Normal"]
    body.fontSize = 11
    body.leading = 15
    heading = styles["Heading2"]
    story = []
    lines = (content or "").splitlines() or [""]
    for line in lines:
        text = escape(line) if line.strip() else "&nbsp;"
        if line.startswith("# "):
            story.append(Paragraph(escape(line[2:].strip()), styles["Heading1"]))
        elif line.startswith("## "):
            story.append(Paragraph(escape(line[3:].strip()), heading))
        else:
            story.append(Paragraph(text, body))
        story.append(Spacer(1, 4))
    template.build(story)
    return buf.getvalue()
